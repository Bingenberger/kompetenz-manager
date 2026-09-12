import os
import secrets
import mimetypes
import re
from datetime import date, datetime, timedelta

from openpyxl import load_workbook
from flask import Blueprint, abort, flash, redirect, render_template, request, send_file, url_for
from flask_login import current_user, login_required
from sqlalchemy import func
from werkzeug.security import generate_password_hash

from extensions import db
from models import Bogen, Beobachtung, Elternkontakt, ErziehungsEreignis, ErziehungsEreignisAnhang, Foerderplan, Item, Notification, Schueler, SystemKonfiguration, User, WorkPlan, WorkPlanTaskAttachment
from odt_export import build_odt_document, convert_odt_bytes_to_pdf
from school_year import observation_period_start
from search import search as run_search
from student_record import collect_record, filename_stem, record_blocks
from student_selection import (
    get_grouped_student_choices_for_user,
    get_prioritized_students_for_user,
    get_tabbed_student_selection_for_user,
    get_user_klassenkontext,
)
from time_utils import utc_now
from uploads import resolve_existing_upload_path

system_bp = Blueprint('system', __name__)


# Reihenfolge zaehlt. Ein Datum wie "2018-07-04" ist eindeutig jahresfuehrend,
# "04.07.2018" eindeutig tagesfuehrend - wer beides mit derselben Annahme liest,
# vertauscht in einem der Faelle Tag und Monat.
_DATE_FORMATS_YEAR_FIRST = ('%Y-%m-%d', '%Y/%m/%d', '%Y.%m.%d')
_DATE_FORMATS_DAY_FIRST = ('%d.%m.%Y', '%d.%m.%y', '%d/%m/%Y', '%d-%m-%Y')


def _parse_import_date(value):
    """Liest ein Geburtsdatum aus einer Excel-Zelle.

    Richtig formatierte Datumszellen liefert openpyxl bereits als datetime.
    Text wird nach Format unterschieden: vierstelliges Jahr vorn heisst
    jahresfuehrend, sonst tagesfuehrend. Was sich nicht lesen laesst, ergibt
    None - ein unleserliches Datum darf nicht die ganze Klassenliste verhindern.
    """
    if value is None:
        return None

    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value

    text_value = str(value).strip()
    if not text_value:
        return None

    formate = (
        _DATE_FORMATS_YEAR_FIRST
        if re.match(r'^\d{4}[-/.]', text_value)
        else _DATE_FORMATS_DAY_FIRST
    )
    for format_string in formate:
        try:
            return datetime.strptime(text_value, format_string).date()
        except ValueError:
            continue
    return None


def _read_xlsx_rows(file_storage):
    """Liest die erste Tabelle einer .xlsx-Datei als (Spalten, Zeilen).

    Zeilen sind dicts mit den Spaltenueberschriften als Schluessel. read_only
    haelt den Speicherbedarf klein - auf einem Schulserver mit knappem RAM ist
    das der Unterschied zwischen Import und Absturz.
    """
    workbook = load_workbook(file_storage, read_only=True, data_only=True)
    try:
        blatt = workbook.active
        zeilen_iterator = blatt.iter_rows(values_only=True)

        try:
            kopf = next(zeilen_iterator)
        except StopIteration:
            return [], []

        spalten = [
            str(zelle).strip() if zelle is not None else ''
            for zelle in kopf
        ]

        zeilen = []
        for werte in zeilen_iterator:
            # Tabellen aus der Praxis enden oft mit leeren Zeilen.
            if all(wert is None or str(wert).strip() == '' for wert in werte):
                continue
            zeile = {}
            for spalte, wert in zip(spalten, werte):
                if spalte:
                    zeile[spalte] = wert
            zeilen.append(zeile)

        return [spalte for spalte in spalten if spalte], zeilen
    finally:
        workbook.close()


def _import_text(value):
    """Zellinhalt als Text - Excel liefert eine Klasse "4" als Zahl."""
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _get_system_konfiguration():
    return SystemKonfiguration.query.first()


def _get_dashboard_fokus_klasse(user):
    klassenkontext = get_user_klassenkontext(user)
    klassenleitung = klassenkontext.get('klassenleitung')
    fachklassen = sorted(klassenkontext.get('fachklassen') or set(), key=lambda x: x.lower())
    fokus_klasse = klassenleitung or (fachklassen[0] if fachklassen else None)
    fokus_typ = 'klassenleitung' if klassenleitung else ('fach' if fokus_klasse else None)
    return fokus_klasse, fokus_typ


def _get_elternsprechtag_fenster(config, heute):
    termine = []
    if config:
        if config.elternsprechtag_1:
            tage = (config.elternsprechtag_1 - heute).days
            termine.append(("Elternsprechtag 1", config.elternsprechtag_1, tage))
        if config.elternsprechtag_2:
            tage = (config.elternsprechtag_2 - heute).days
            termine.append(("Elternsprechtag 2", config.elternsprechtag_2, tage))
    return termine


def _todo_sort_key(todo):
    priority = int(todo.get('priority', 99))
    due_date = todo.get('due_date')
    due_ord = due_date.toordinal() if due_date else 999999999
    title = (todo.get('title') or '').lower()
    return (priority, due_ord, title)


def _get_notification_or_404(notification_id):
    notification = db.session.get(Notification, notification_id)
    if not notification:
        abort(404)
    if notification.user_id != current_user.id and not current_user.is_admin:
        abort(403)
    return notification


def _is_admin(user):
    return bool(user and getattr(user, 'is_admin', False))


def _send_upload_or_404(rel_path, download_name=None, mimetype=None):
    resolved_path = resolve_existing_upload_path(rel_path)
    if not resolved_path:
        abort(404)
    guessed_type, _ = mimetypes.guess_type(str(resolved_path))
    return send_file(
        resolved_path,
        mimetype=mimetype or guessed_type or 'application/octet-stream',
        download_name=download_name,
        conditional=True,
    )

@system_bp.route('/')
@login_required
def index():
    klassenkontext = get_user_klassenkontext(current_user)
    klassenleitung = klassenkontext.get('klassenleitung')
    fachklassen = sorted(klassenkontext.get('fachklassen') or set(), key=lambda x: x.lower())

    fokus_klasse = klassenleitung or (fachklassen[0] if fachklassen else None)
    fokus_typ = 'klassenleitung' if klassenleitung else ('fach' if fokus_klasse else None)

    stats = {
        'kinder_in_klasse': 0,
        'beobachtungsboegen_ausgefuellt': 0,
        'aktive_foerderplaene': 0,
        'aktive_arbeitsplaene': 0,
        'offene_erziehungsfaelle': 0,
    }

    todos = []
    heute = utc_now().date()
    in_14_tagen = heute + timedelta(days=14)
    in_12_wochen_datetime = observation_period_start(utc_now() - timedelta(weeks=12))
    config = _get_system_konfiguration()

    if fokus_klasse:
        klasse_kinder = (
            Schueler.query
            .filter(Schueler.klasse == fokus_klasse, Schueler.is_active.is_(True))
            .order_by(Schueler.nachname.asc(), Schueler.vorname.asc())
            .all()
        )
        stats['kinder_in_klasse'] = len(klasse_kinder)

        stats['beobachtungsboegen_ausgefuellt'] = (
            Beobachtung.query
            .join(Schueler, Beobachtung.schueler_id == Schueler.id)
            .join(Item, Beobachtung.item_id == Item.id)
            .filter(Schueler.klasse == fokus_klasse, Schueler.is_active.is_(True))
            .with_entities(Beobachtung.schueler_id, Item.bogen_id)
            .distinct()
            .count()
        )

        stats['aktive_foerderplaene'] = (
            Foerderplan.query
            .join(Schueler, Foerderplan.schueler_id == Schueler.id)
            .filter(
                Schueler.klasse == fokus_klasse,
                    Schueler.is_active.is_(True),
                Foerderplan.status == 'aktiv',
            )
            .count()
        )

        stats['aktive_arbeitsplaene'] = (
            WorkPlan.query
            .join(Schueler, WorkPlan.student_id == Schueler.id)
            .filter(
                Schueler.klasse == fokus_klasse,
                    Schueler.is_active.is_(True),
                WorkPlan.status.in_(['draft', 'active']),
            )
            .count()
        )

        stats['offene_erziehungsfaelle'] = (
            ErziehungsEreignis.query
            .join(Schueler, ErziehungsEreignis.student_id == Schueler.id)
            .filter(
                Schueler.klasse == fokus_klasse,
                    Schueler.is_active.is_(True),
                ErziehungsEreignis.status == 'offen',
            )
            .count()
        )

        aktive_plan_schueler_ids = {
            row[0]
            for row in (
                Foerderplan.query
                .join(Schueler, Foerderplan.schueler_id == Schueler.id)
                .filter(
                    Schueler.klasse == fokus_klasse,
                    Schueler.is_active.is_(True),
                    Foerderplan.status == 'aktiv',
                )
                .with_entities(Foerderplan.schueler_id)
                .distinct()
                .all()
            )
        }

        elternsprechtage = []
        if config:
            if config.elternsprechtag_1:
                elternsprechtage.append(("Elternsprechtag 1", config.elternsprechtag_1))
            if config.elternsprechtag_2:
                elternsprechtage.append(("Elternsprechtag 2", config.elternsprechtag_2))

        hat_4_wochen_fenster = False
        hat_2_wochen_fenster = False
        aktive_esp_termine = []
        for label, termin in elternsprechtage:
            tage_bis = (termin - heute).days
            if 1 <= tage_bis <= 28:
                hat_4_wochen_fenster = True
                aktive_esp_termine.append((label, termin, tage_bis))
            if 1 <= tage_bis <= 14:
                hat_2_wochen_fenster = True

        if hat_4_wochen_fenster:
            for label, termin, _tage_bis in aktive_esp_termine:
                if stats['aktive_foerderplaene'] > 0:
                    todos.append({
                        'title': f'{label}: aktive Förderpläne evaluieren',
                        'detail': (
                            f"{stats['aktive_foerderplaene']} aktive Förderpläne in Klasse {fokus_klasse} "
                            f"vor dem Termin am {termin.strftime('%d.%m.%Y')} prüfen."
                        ),
                        'variant': 'warning',
                        'url': url_for('system.todo_foerderplan_evaluationen'),
                        'priority': 30,
                        'due_date': termin,
                    })

            minus_rows = (
                db.session.query(
                    Beobachtung.schueler_id.label('schueler_id'),
                    Item.bogen_id.label('bogen_id'),
                    func.count(func.distinct(Beobachtung.item_id)).label('minus_items'),
                )
                .join(Schueler, Beobachtung.schueler_id == Schueler.id)
                .join(Item, Beobachtung.item_id == Item.id)
                .filter(
                    Schueler.klasse == fokus_klasse,
                    Schueler.is_active.is_(True),
                    Beobachtung.wert == 1,
                    Beobachtung.datum >= in_12_wochen_datetime,
                )
                .group_by(Beobachtung.schueler_id, Item.bogen_id)
                .all()
            )

            foerderplan_kandidaten_ids = {
                row.schueler_id
                for row in minus_rows
                if (row.minus_items or 0) >= 4 and row.schueler_id not in aktive_plan_schueler_ids
            }

            if foerderplan_kandidaten_ids:
                kandidaten = [s for s in klasse_kinder if s.id in foerderplan_kandidaten_ids]
                beispiele = ', '.join(f'{s.vorname} {s.nachname}'.strip() for s in kandidaten[:3])
                beispiel_text = f' Beispiele: {beispiele}.' if beispiele else ''
                for label, termin, _tage_bis in aktive_esp_termine:
                    todos.append({
                        'title': f'{label}: Förderplan-Kandidaten prüfen',
                        'detail': (
                            f"{len(foerderplan_kandidaten_ids)} Kind(er) mit mind. 4x '-' in einem Bogen "
                            f"(letzte 12 Wochen) ohne aktiven Förderplan. Termin: {termin.strftime('%d.%m.%Y')}."
                            f"{beispiel_text}"
                        ),
                        'variant': 'danger',
                        'url': url_for('system.todo_foerderplan_kandidaten'),
                        'priority': 20,
                        'due_date': termin,
                    })

        if hat_2_wochen_fenster and klasse_kinder:
            schueler_mit_aktuellem_bogen_ids = {
                row[0]
                for row in (
                    Beobachtung.query
                    .join(Schueler, Beobachtung.schueler_id == Schueler.id)
                    .filter(
                        Schueler.klasse == fokus_klasse,
                    Schueler.is_active.is_(True),
                        Beobachtung.datum >= in_12_wochen_datetime,
                    )
                    .with_entities(Beobachtung.schueler_id)
                    .distinct()
                    .all()
                )
            }
            fehlende_aktuelle_boegen = [
                s for s in klasse_kinder if s.id not in schueler_mit_aktuellem_bogen_ids
            ]
            if fehlende_aktuelle_boegen:
                naechster_termin_2w = min(
                    (termin for _label, termin, tage in aktive_esp_termine if 1 <= tage <= 14),
                    default=None,
                )
                beispiele = ', '.join(
                    f'{s.vorname} {s.nachname}'.strip() for s in fehlende_aktuelle_boegen[:3]
                )
                todos.append({
                    'title': 'Beobachtungsbögen aktualisieren',
                    'detail': (
                        f"{len(fehlende_aktuelle_boegen)} Kind(er) ohne Beobachtungsbogen-Eintrag in den letzten 12 Wochen"
                        + (f" vor dem Elternsprechtag am {naechster_termin_2w.strftime('%d.%m.%Y')}" if naechster_termin_2w else '')
                        + (f". Beispiele: {beispiele}." if beispiele else '.')
                    ),
                    'variant': 'info',
                    'url': url_for('system.todo_beobachtungsboegen_fehlend'),
                    'priority': 40,
                    'due_date': naechster_termin_2w,
                })

        ueberfaellige_evals = (
            Foerderplan.query
            .join(Schueler, Foerderplan.schueler_id == Schueler.id)
            .filter(
                Schueler.klasse == fokus_klasse,
                    Schueler.is_active.is_(True),
                Foerderplan.status == 'aktiv',
                Foerderplan.datum_evaluation.isnot(None),
                Foerderplan.datum_evaluation < heute,
            )
            .count()
        )
        if ueberfaellige_evals:
            todos.append({
                'title': f'{ueberfaellige_evals} Förderplan-Evaluation(en) überfällig',
                'detail': f'Klasse {fokus_klasse} bitte priorisiert prüfen.',
                'variant': 'danger',
                'url': url_for('system.todo_foerderplan_evaluationen'),
                'priority': 10,
                'due_date': heute,
            })

        anstehende_evals = (
            Foerderplan.query
            .join(Schueler, Foerderplan.schueler_id == Schueler.id)
            .filter(
                Schueler.klasse == fokus_klasse,
                    Schueler.is_active.is_(True),
                Foerderplan.status == 'aktiv',
                Foerderplan.datum_evaluation.isnot(None),
                Foerderplan.datum_evaluation >= heute,
                Foerderplan.datum_evaluation <= in_14_tagen,
            )
            .order_by(Foerderplan.datum_evaluation.asc())
            .limit(3)
            .all()
        )
        for plan in anstehende_evals:
            schueler_name = f"{plan.schueler.vorname or ''} {plan.schueler.nachname or ''}".strip() or 'Kind'
            todos.append({
                'title': f'Evaluation Förderplan: {schueler_name}',
                'detail': f"Fällig bis {plan.datum_evaluation.strftime('%d.%m.%Y')}",
                'variant': 'warning',
                'url': url_for('system.todo_foerderplan_evaluationen'),
                'priority': 25,
                'due_date': plan.datum_evaluation,
            })

        naechste_termine = (
            Elternkontakt.query
            .join(Schueler, Elternkontakt.schueler_id == Schueler.id)
            .filter(
                Schueler.klasse == fokus_klasse,
                    Schueler.is_active.is_(True),
                Elternkontakt.naechster_termin.isnot(None),
                Elternkontakt.naechster_termin >= heute,
                Elternkontakt.naechster_termin <= in_14_tagen,
            )
            .order_by(Elternkontakt.naechster_termin.asc())
            .limit(2)
            .all()
        )
        for kontakt in naechste_termine:
            schueler_name = f"{kontakt.schueler.vorname or ''} {kontakt.schueler.nachname or ''}".strip() or 'Kind'
            todos.append({
                'title': f'Elternkontakt nachverfolgen: {schueler_name}',
                'detail': f"Termin am {kontakt.naechster_termin.strftime('%d.%m.%Y')}",
                'variant': 'info',
                'url': url_for('system.todo_elternkontakt_erinnerungen'),
                'priority': 35,
                'due_date': kontakt.naechster_termin,
            })

        offene_faelle = (
            ErziehungsEreignis.query
            .join(Schueler, ErziehungsEreignis.student_id == Schueler.id)
            .filter(
                Schueler.klasse == fokus_klasse,
                    Schueler.is_active.is_(True),
                ErziehungsEreignis.status == 'offen',
            )
            .order_by(ErziehungsEreignis.datum.asc(), ErziehungsEreignis.id.asc())
            .limit(3)
            .all()
        )
        for fall in offene_faelle:
            schueler_name = f"{fall.student.vorname or ''} {fall.student.nachname or ''}".strip() or 'Kind'
            ereignis_name = fall.event_template.name if fall.event_template else 'Ereignis'
            zustaendig = fall.assigned_user.display_name if fall.assigned_user else 'noch nicht zugewiesen'
            todos.append({
                'title': f'Offener Fall: {schueler_name}',
                'detail': f"{ereignis_name} vom {fall.datum.strftime('%d.%m.%Y')} · Zuständig: {zustaendig}",
                'variant': 'danger' if not fall.assigned_user_id else 'warning',
                'url': url_for('erziehung.erziehung_view', event_id=fall.id, next=url_for('system.index')),
                'priority': 15 if not fall.assigned_user_id else 22,
                'due_date': fall.datum,
            })

        if stats['beobachtungsboegen_ausgefuellt'] == 0:
            todos.append({
                'title': 'Erste Beobachtungsbögen erfassen',
                'detail': f'Für Klasse {fokus_klasse} liegen noch keine ausgefüllten Bögen vor.',
                'variant': 'success',
                'url': url_for('erfassung.erfassen_schueler'),
                'priority': 50,
            })
    else:
        todos.append({
            'title': 'Klasse zuordnen lassen',
            'detail': 'Für personalisierte Statistik bitte eine Klassenleitung oder Fachklasse hinterlegen.',
            'variant': 'info',
            'url': url_for('auth.user_menu'),
            'priority': 5,
        })

    if not todos:
        todos.append({
            'title': 'Aktuell keine dringenden Aufgaben',
            'detail': 'In den nächsten 14 Tagen sind keine Termine oder Evaluationen hinterlegt.',
            'variant': 'success',
            'url': url_for('system.schuelerakte'),
            'priority': 99,
        })

    todos.sort(key=_todo_sort_key)

    return render_template(
        'index.html',
        teacher_name=current_user.display_name,
        hero_klasse=fokus_klasse,
        hero_klasse_typ=fokus_typ,
        hero_stats=stats,
        hero_todos=todos[:8],
        school_year=(config.schuljahr if config else None),
    )


@system_bp.route('/suche')
@login_required
def suche():
    query = (request.args.get('q') or '').strip()
    ergebnis = run_search(current_user, query)

    # Ein einzelner Treffer braucht keine Ergebnisliste: wer "abt" eintippt und
    # genau ein Kind findet, will in die Schuelerakte, nicht auf eine Seite mit
    # einem Link darauf.
    if ergebnis['total'] == 1 and len(ergebnis['students']) == 1:
        return redirect(url_for('system.schuelerakte', schueler_id=ergebnis['students'][0].id))

    return render_template('suche.html', ergebnis=ergebnis)


@system_bp.route('/schuelerakte')
@login_required
def schuelerakte():
    selection = get_tabbed_student_selection_for_user(
        current_user,
        selected_s_id=(request.args.get('schueler_id') or '').strip(),
        requested_tab=(request.args.get('tab') or '').strip(),
        auto_select_first=True,
        include_archived=True,
    )
    schueler_liste = selection['students']
    schueler_groups = selection['groups']
    selected_s_id = selection['selected_s_id']
    selected_student = selection['selected_student']

    foerderplaene = []
    work_plans = []
    erziehungsereignisse = []
    elternkontakte = []
    bogen_summaries = []
    recent_beobachtungen = []

    if selected_student:
        foerderplaene = (
            Foerderplan.query
            .filter(Foerderplan.schueler_id == selected_student.id)
            .order_by(Foerderplan.datum_erstellung.desc())
            .all()
        )
        work_plans = (
            WorkPlan.query
            .filter(WorkPlan.student_id == selected_student.id)
            .order_by(WorkPlan.created_at.desc())
            .all()
        )
        erziehungsereignisse = (
            ErziehungsEreignis.query
            .filter(ErziehungsEreignis.student_id == selected_student.id)
            .order_by(ErziehungsEreignis.datum.desc(), ErziehungsEreignis.id.desc())
            .all()
        )

        elternkontakte = (
            Elternkontakt.query
            .filter(Elternkontakt.schueler_id == selected_student.id)
            .order_by(Elternkontakt.datum.desc())
            .limit(12)
            .all()
        )

        bogen_rows = (
            db.session.query(
                Bogen.id.label('bogen_id'),
                Bogen.titel.label('bogen_titel'),
                func.count(Beobachtung.id).label('anzahl_eintraege'),
                func.max(Beobachtung.datum).label('letzte_beobachtung'),
            )
            .select_from(Beobachtung)
            .join(Item, Beobachtung.item_id == Item.id)
            .join(Bogen, Item.bogen_id == Bogen.id)
            .filter(Beobachtung.schueler_id == selected_student.id)
            .group_by(Bogen.id, Bogen.titel)
            .order_by(func.max(Beobachtung.datum).desc())
            .all()
        )
        bogen_summaries = list(bogen_rows)

        beobachtungs_rows = (
            db.session.query(Beobachtung, Item.text, Bogen.titel)
            .join(Item, Beobachtung.item_id == Item.id)
            .join(Bogen, Item.bogen_id == Bogen.id)
            .filter(Beobachtung.schueler_id == selected_student.id)
            .order_by(Beobachtung.datum.desc())
            .limit(12)
            .all()
        )
        for beobachtung, item_text, bogen_titel in beobachtungs_rows:
            recent_beobachtungen.append({
                'beobachtung': beobachtung,
                'item_text': item_text,
                'bogen_titel': bogen_titel,
            })

    return render_template(
        'schuelerakte.html',
        schueler=schueler_liste,
        schueler_groups=schueler_groups,
        tab_definitions=selection['tab_definitions'],
        active_tab=selection['active_tab'],
        selected_s_id=selected_s_id,
        selected_student=selected_student,
        foerderplaene=foerderplaene,
        work_plans=work_plans,
        erziehungsereignisse=erziehungsereignisse,
        elternkontakte=elternkontakte,
        bogen_summaries=bogen_summaries,
        recent_beobachtungen=recent_beobachtungen,
    )


def _build_student_record_document(schueler):
    """Baut das ODT der vollstaendigen Akte und liefert Puffer samt Dateiname."""
    jetzt = utc_now()
    record = collect_record(schueler)
    blocks = record_blocks(record, jetzt.strftime('%d.%m.%Y'))
    return build_odt_document(blocks), filename_stem(schueler, jetzt)


@system_bp.route('/schuelerakte/export/odt/<int:s_id>')
@login_required
def schuelerakte_export_odt(s_id):
    schueler = db.session.get(Schueler, s_id)
    if not schueler:
        abort(404)

    odt_buffer, stem = _build_student_record_document(schueler)
    return send_file(
        odt_buffer,
        as_attachment=True,
        download_name=f'{stem}.odt',
        mimetype='application/vnd.oasis.opendocument.text',
    )


@system_bp.route('/schuelerakte/export/pdf/<int:s_id>')
@login_required
def schuelerakte_export_pdf(s_id):
    schueler = db.session.get(Schueler, s_id)
    if not schueler:
        abort(404)

    next_url = url_for('system.schuelerakte', schueler_id=s_id)
    odt_buffer, stem = _build_student_record_document(schueler)
    try:
        pdf_buffer = convert_odt_bytes_to_pdf(odt_buffer)
    except RuntimeError as exc:
        # LibreOffice fehlt oder bricht ab - das ODT bleibt der Weg.
        flash(f'PDF-Export fehlgeschlagen: {exc}. Der ODT-Export funktioniert unabhängig davon.')
        return redirect(next_url)

    return send_file(
        pdf_buffer,
        as_attachment=True,
        download_name=f'{stem}.pdf',
        mimetype='application/pdf',
    )


@system_bp.route('/todo/elternkontakte')
@login_required
def todo_elternkontakt_erinnerungen():
    heute = utc_now().date()
    in_14_tagen = heute + timedelta(days=14)
    fokus_klasse, _ = _get_dashboard_fokus_klasse(current_user)
    if not fokus_klasse:
        flash('Keine Klasse zugeordnet.')
        return redirect(url_for('system.index'))

    termine = (
        Elternkontakt.query
        .join(Schueler, Elternkontakt.schueler_id == Schueler.id)
        .filter(
            Schueler.klasse == fokus_klasse,
                    Schueler.is_active.is_(True),
            Elternkontakt.naechster_termin.isnot(None),
            Elternkontakt.naechster_termin >= heute,
            Elternkontakt.naechster_termin <= in_14_tagen,
        )
        .order_by(Elternkontakt.naechster_termin.asc())
        .all()
    )

    rows = []
    for kontakt in termine:
        last_protocol = (
            Elternkontakt.query
            .filter(
                Elternkontakt.schueler_id == kontakt.schueler_id,
                Elternkontakt.eintrag_typ == 'protokoll',
            )
            .order_by(Elternkontakt.datum.desc())
            .first()
        )
        rows.append({
            'kontakt': kontakt,
            'last_protocol': last_protocol,
        })

    return render_template(
        'todo_elternkontakte.html',
        fokus_klasse=fokus_klasse,
        rows=rows,
        heute=heute,
    )


@system_bp.route('/todo/foerderplan-evaluationen')
@login_required
def todo_foerderplan_evaluationen():
    heute = utc_now().date()
    fokus_klasse, _ = _get_dashboard_fokus_klasse(current_user)
    if not fokus_klasse:
        flash('Keine Klasse zugeordnet.')
        return redirect(url_for('system.index'))

    config = _get_system_konfiguration()
    termine = _get_elternsprechtag_fenster(config, heute)
    aktive_4w = [t for t in termine if 1 <= t[2] <= 28]

    plaene = (
        Foerderplan.query
        .join(Schueler, Foerderplan.schueler_id == Schueler.id)
        .filter(
            Schueler.klasse == fokus_klasse,
                    Schueler.is_active.is_(True),
            Foerderplan.status == 'aktiv',
        )
        .order_by(Schueler.nachname.asc(), Schueler.vorname.asc(), Foerderplan.datum_erstellung.desc())
        .all()
    )

    return render_template(
        'todo_foerderplan_evaluationen.html',
        fokus_klasse=fokus_klasse,
        plaene=plaene,
        aktive_elternsprechtage=aktive_4w,
        heute=heute,
    )


@system_bp.route('/todo/foerderplan-kandidaten')
@login_required
def todo_foerderplan_kandidaten():
    heute = utc_now().date()
    fokus_klasse, _ = _get_dashboard_fokus_klasse(current_user)
    if not fokus_klasse:
        flash('Keine Klasse zugeordnet.')
        return redirect(url_for('system.index'))

    in_12_wochen_datetime = observation_period_start(utc_now() - timedelta(weeks=12))
    config = _get_system_konfiguration()
    aktive_4w = [t for t in _get_elternsprechtag_fenster(config, heute) if 1 <= t[2] <= 28]

    aktive_plan_schueler_ids = {
        row[0]
        for row in (
            Foerderplan.query
            .join(Schueler, Foerderplan.schueler_id == Schueler.id)
            .filter(Schueler.klasse == fokus_klasse, Schueler.is_active.is_(True), Foerderplan.status == 'aktiv')
            .with_entities(Foerderplan.schueler_id)
            .distinct()
            .all()
        )
    }

    minus_rows = (
        db.session.query(
            Beobachtung.schueler_id.label('schueler_id'),
            Item.bogen_id.label('bogen_id'),
            func.count(func.distinct(Beobachtung.item_id)).label('minus_items'),
        )
        .join(Schueler, Beobachtung.schueler_id == Schueler.id)
        .join(Item, Beobachtung.item_id == Item.id)
        .filter(
            Schueler.klasse == fokus_klasse,
                    Schueler.is_active.is_(True),
            Beobachtung.wert == 1,
            Beobachtung.datum >= in_12_wochen_datetime,
        )
        .group_by(Beobachtung.schueler_id, Item.bogen_id)
        .all()
    )

    bogens = {b.id: b for b in Bogen.query.all()}
    schueler_map = {
        s.id: s for s in Schueler.query.filter(Schueler.klasse == fokus_klasse, Schueler.is_active.is_(True)).all()
    }

    kandidaten = {}
    for row in minus_rows:
        if (row.minus_items or 0) < 4 or row.schueler_id in aktive_plan_schueler_ids:
            continue
        s = schueler_map.get(row.schueler_id)
        if not s:
            continue
        eintrag = kandidaten.setdefault(s.id, {'schueler': s, 'boegen': []})
        eintrag['boegen'].append({
            'bogen_id': row.bogen_id,
            'bogen_titel': getattr(bogens.get(row.bogen_id), 'titel', f'Bogen {row.bogen_id}'),
            'minus_items': int(row.minus_items or 0),
        })

    kandidat_rows = sorted(
        kandidaten.values(),
        key=lambda r: ((r['schueler'].nachname or '').lower(), (r['schueler'].vorname or '').lower()),
    )

    return render_template(
        'todo_foerderplan_kandidaten.html',
        fokus_klasse=fokus_klasse,
        kandidat_rows=kandidat_rows,
        aktive_elternsprechtage=aktive_4w,
    )


@system_bp.route('/todo/beobachtungsboegen-fehlend')
@login_required
def todo_beobachtungsboegen_fehlend():
    heute = utc_now().date()
    fokus_klasse, _ = _get_dashboard_fokus_klasse(current_user)
    if not fokus_klasse:
        flash('Keine Klasse zugeordnet.')
        return redirect(url_for('system.index'))

    in_12_wochen_datetime = observation_period_start(utc_now() - timedelta(weeks=12))
    config = _get_system_konfiguration()
    aktive_2w = [t for t in _get_elternsprechtag_fenster(config, heute) if 1 <= t[2] <= 14]

    kinder = (
        Schueler.query
        .filter(Schueler.klasse == fokus_klasse, Schueler.is_active.is_(True))
        .order_by(Schueler.nachname.asc(), Schueler.vorname.asc())
        .all()
    )
    boegen = Bogen.query.order_by(Bogen.titel.asc()).all()

    aktuelle_kombis = {
        (row[0], row[1])
        for row in (
            Beobachtung.query
            .join(Schueler, Beobachtung.schueler_id == Schueler.id)
            .join(Item, Beobachtung.item_id == Item.id)
            .filter(
                Schueler.klasse == fokus_klasse,
                    Schueler.is_active.is_(True),
                Beobachtung.datum >= in_12_wochen_datetime,
            )
            .with_entities(Beobachtung.schueler_id, Item.bogen_id)
            .distinct()
            .all()
        )
    }

    rows = []
    for s in kinder:
        fehlende_boegen = []
        for b in boegen:
            if (s.id, b.id) not in aktuelle_kombis:
                fehlende_boegen.append(b)
        if fehlende_boegen:
            rows.append({'schueler': s, 'fehlende_boegen': fehlende_boegen})

    return render_template(
        'todo_beobachtungsboegen_fehlend.html',
        fokus_klasse=fokus_klasse,
        rows=rows,
        aktive_elternsprechtage=aktive_2w,
    )

@system_bp.route('/import/<typ>', methods=['GET', 'POST'])
@login_required
def data_import(typ):
    next_url = (request.args.get('next') or request.form.get('next') or '').strip()
    if request.method == 'POST':
        file = request.files['file']
        if file.filename.endswith('.xlsx'):
            spalten, zeilen = _read_xlsx_rows(file)
            required_columns = {
                'schueler': ['Vorname', 'Nachname', 'Klasse', 'Geburtsdatum'],
                'bogen': ['Bogen', 'Bereich', 'Item'],
            }.get(typ, [])
            missing_columns = [col for col in required_columns if col not in spalten]
            if missing_columns:
                flash('Fehlende Spalten in Excel-Datei: ' + ', '.join(missing_columns))
                return redirect(url_for('system.data_import', typ=typ, next=next_url) if next_url else url_for('system.data_import', typ=typ))
            if typ == 'schueler':
                for row in zeilen:
                    db.session.add(Schueler(
                        vorname=_import_text(row.get('Vorname')),
                        nachname=_import_text(row.get('Nachname')),
                        klasse=_import_text(row.get('Klasse')),
                        geburtsdatum=_parse_import_date(row.get('Geburtsdatum')),
                    ))
            elif typ == 'bogen':
                for row in zeilen:
                    bogen_titel = _import_text(row.get('Bogen'))
                    exist_bogen = Bogen.query.filter_by(titel=bogen_titel).first()
                    if not exist_bogen:
                        exist_bogen = Bogen(titel=bogen_titel)
                        db.session.add(exist_bogen)
                        db.session.flush()

                    item = Item(
                        bogen_id=exist_bogen.id,
                        bereich=_import_text(row.get('Bereich')),
                        text=_import_text(row.get('Item')),
                    )
                    db.session.add(item)

            db.session.commit()
            flash(f'{typ.capitalize()} erfolgreich importiert!')
            if next_url.startswith('/') and not next_url.startswith('//'):
                return redirect(next_url)
            return redirect(url_for('system.index'))
    return render_template('import.html', typ=typ, next_url=next_url)

@system_bp.route('/setup')
def setup():
    if os.environ.get('ALLOW_SETUP_ROUTE') != '1':
        return "Setup-Route ist deaktiviert.", 403

    setup_token = (os.environ.get('SETUP_ROUTE_TOKEN') or '').strip()
    if not setup_token:
        return "Setup-Token fehlt. Route bleibt deaktiviert.", 403

    request_token = (request.args.get('token') or '').strip()
    if not request_token or not secrets.compare_digest(request_token, setup_token):
        return "Ungueltiger Setup-Token.", 403

    if request.remote_addr not in {'127.0.0.1', '::1'}:
        return "Setup ist nur lokal auf dem Server erlaubt.", 403

    db.create_all()
    if not User.query.filter_by(role='admin').first():
        setup_password = os.environ.get('SETUP_ADMIN_PASSWORD') or secrets.token_urlsafe(12)
        hashed_pw = generate_password_hash(setup_password)
        admin = User(username='admin', password_hash=hashed_pw, role='admin')
        db.session.add(admin)
        db.session.commit()
        return (
            "Datenbank & Admin-User angelegt! Temporäres Passwort: "
            f"<strong>{setup_password}</strong> "
            f"<a href='{url_for('auth.login')}'>Zum Login</a>"
        )

    if not Schueler.query.first():
        db.session.add(Schueler(vorname="Max", nachname="Muster", klasse="4a"))
        db.session.add(Schueler(vorname="Lisa", nachname="Lustig", klasse="4a"))

        b = Bogen(titel="Sozialverhalten")
        db.session.add(b)
        db.session.flush()
        db.session.add(Item(bogen_id=b.id, bereich="Konflikt", text="Löst Streit friedlich"))
        db.session.add(Item(bogen_id=b.id, bereich="Arbeit", text="Arbeitet konzentriert"))

        db.session.commit()
        return f"Datenbank erstellt und Testdaten angelegt! <a href='{url_for('system.index')}'>Zum Start</a>"

    return f"Datenbank existiert schon. <a href='{url_for('system.index')}'>Zum Start</a>"


@system_bp.route('/media/beobachtung/<int:beobachtung_id>')
@login_required
def media_beobachtung(beobachtung_id):
    observation = db.session.get(Beobachtung, beobachtung_id)
    if not observation or not observation.foto_pfad:
        abort(404)
    return _send_upload_or_404(
        observation.foto_pfad,
        download_name=f'beobachtung-{beobachtung_id}.jpg',
    )


@system_bp.route('/media/workplan-attachment/<string:attachment_id>')
@login_required
def media_workplan_attachment(attachment_id):
    attachment = db.session.get(WorkPlanTaskAttachment, attachment_id)
    if not attachment:
        abort(404)

    task = attachment.task
    plan = task.work_plan if task else None
    if not plan:
        abort(404)

    if not _is_admin(current_user) and plan.created_by_user_id != current_user.id:
        abort(403)

    return _send_upload_or_404(
        attachment.file_path,
        download_name=attachment.caption or f'arbeitsplan-{attachment.id}.jpg',
    )


@system_bp.route('/media/erziehung-attachment/<int:attachment_id>')
@login_required
def media_erziehung_attachment(attachment_id):
    attachment = db.session.get(ErziehungsEreignisAnhang, attachment_id)
    if not attachment:
        abort(404)

    event = attachment.event
    if not event:
        abort(404)

    if not _is_admin(current_user):
        accessible_ids = {student.id for student in get_prioritized_students_for_user(current_user)}
        if event.student_id not in accessible_ids:
            abort(403)

    return _send_upload_or_404(
        attachment.file_path,
        download_name=attachment.original_name or attachment.file_path.rsplit('/', 1)[-1],
        mimetype=attachment.mime_type or None,
    )


@system_bp.route('/benachrichtigungen/<int:notification_id>/open', methods=['POST'])
@login_required
def notification_open(notification_id):
    notification = _get_notification_or_404(notification_id)
    notification.is_read = True
    db.session.commit()
    return redirect(notification.target_url or url_for('system.index'))


@system_bp.route('/benachrichtigungen/<int:notification_id>/delete', methods=['POST'])
@login_required
def notification_delete(notification_id):
    notification = _get_notification_or_404(notification_id)
    db.session.delete(notification)
    db.session.commit()
    flash('Benachrichtigung gelöscht.')
    return redirect((request.form.get('next') or '').strip() or url_for('system.index'))


@system_bp.route('/benachrichtigungen/delete-all', methods=['POST'])
@login_required
def notification_delete_all():
    Notification.query.filter_by(user_id=current_user.id).delete()
    db.session.commit()
    flash('Benachrichtigungen gelöscht.')
    return redirect((request.form.get('next') or '').strip() or url_for('system.index'))


def register_system_routes(app):
    app.register_blueprint(system_bp)

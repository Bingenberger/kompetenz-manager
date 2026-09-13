"""Tests fuer den Import aus den Auswertungsmappen (HSP, ELFE II, SLS).

Die Testdateien werden hier im Aufbau der Schulvorlagen erzeugt - mit
erfundenen Namen. Die echten Mappen liegen nur lokal unter import/ und
gehoeren nicht ins Repository.
"""

import os
import re
import shutil
import tempfile
import unittest
import zipfile
from datetime import date
from io import BytesIO
from xml.sax.saxutils import escape

from openpyxl import Workbook
from werkzeug.security import generate_password_hash

from app import create_app
from diagnostik import lege_vorbelegung_an, schaerfe_vorbelegung_nach
from diagnostik_import import ImportFehler, lese_import, normalisiere_name, ordne_kinder_zu, zahl
from extensions import db
from jahrgang import ensure_klasse
from models import (
    DiagnostikErgebnis,
    DiagnostikKennwert,
    DiagnostikTestform,
    DiagnostikWert,
    DiagnostikZeitpunkt,
    Schueler,
    SystemKonfiguration,
    User,
    UserKlassenzuordnung,
)

CSRF_RE = re.compile(r'name="_csrf_token"\s+value="([^"]+)"')


# ----------------------------------------------------------------------
# Testdateien im Aufbau der Schulvorlagen
# ----------------------------------------------------------------------

def hsp_mappe():
    mappe = Workbook()
    blatt = mappe.active
    blatt.title = 'Gesamt'
    blatt['B1'] = 'Anzahl Kinder'
    blatt['D1'] = 30
    kopf = ['Nr.', 'Vorname', 'Name', 'W', 'PR', 'T', 'GT', 'PR2', 'T3', 'AS', 'PR4', 'T5',
            'OS', 'PR6', 'T7', 'MS', 'PR8', 'T9']
    for spalte, wert in enumerate(kopf, start=2):
        blatt.cell(3, spalte, wert)
    zeilen = [
        [1, 'Anna', 'Abt', 25, 40, 47, 180, 35, 46, 18, 50, 50, 12, 30, 45, 8, 20, 42],
        [2, 'Ben', 'Bär', 10, 5, 34, 120, 8, 36, 12, 10, 37, 6, 7, 35, 3, 4, 33],
        [3, 'VORNAME', 'NACHNAME', 38, 99, 72, 191, 99, 72, 20, 75, 57, 15, 85, 60, 10, 90, 63],
        [4, 'Zora', 'Zett', 20, 30, 45, 150, 25, 43, 15, 30, 45, 10, 20, 42, 5, 10, 37],
        [99, 'Schnitt', None, 20, 30, 45, 150, 25, 43, 15, 30, 45, 10, 20, 42, 5, 10, 37],
    ]
    for index, werte in enumerate(zeilen, start=4):
        for spalte, wert in enumerate(werte, start=2):
            blatt.cell(index, spalte, wert)
    mappe.create_sheet('S01')['A1'] = 'Name:'
    puffer = BytesIO()
    mappe.save(puffer)
    return puffer.getvalue()


def elfe_mappe():
    mappe = Workbook()
    blatt = mappe.active
    blatt.title = 'Daten'
    blatt['B1'], blatt['B2'], blatt['B3'], blatt['B4'] = 'Schuljahr', 3, 'Schulmonat', '11-12'
    blatt['A2'] = 'Testzeitpunkt'
    for spalte, gruppe in (('C', 'Wortverständnis'), ('F', 'Satzverständnis'), ('I', 'Textverständnis'), ('L', 'Gesamtauswertung')):
        blatt[f'{spalte}6'] = gruppe
    for spalte, kopf in zip('ABCDEFGHIJKLMN', ['Name', 'Vorname', 'RW', 'T-W', 'PR', 'RW', 'T-W', 'PR', 'RW', 'T-W', 'PR', 'UT', 'T-W', 'PR']):
        blatt[f'{spalte}7'] = kopf
    for spalte, wert in zip('ABCDEFGHIJKLMN', ['Abt', 'Anna', 40, 45, 30.9, 20, 38, 11.5, 10, 30, 2.3, 113, 36, 0.6]):
        blatt[f'{spalte}8'] = wert
    puffer = BytesIO()
    mappe.save(puffer)
    return puffer.getvalue()


def ods_datei(zeilen, blattname='Auswertung SLS'):
    """Minimale ODS-Datei; Zellen mit float werden als Zahl gespeichert."""
    def zelle(wert):
        if wert is None:
            return '<table:table-cell/>'
        if isinstance(wert, (int, float)):
            return f'<table:table-cell office:value-type="float" office:value="{wert}"><text:p>{wert}</text:p></table:table-cell>'
        return f'<table:table-cell office:value-type="string"><text:p>{escape(str(wert))}</text:p></table:table-cell>'

    xml_zeilen = ''.join(
        '<table:table-row>' + ''.join(zelle(w) for w in zeile)
        + '<table:table-cell table:number-columns-repeated="1000"/></table:table-row>'
        for zeile in zeilen
    )
    inhalt = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
        'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
        'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" office:version="1.2">'
        f'<office:body><office:spreadsheet><table:table table:name="{blattname}">{xml_zeilen}'
        '<table:table-row table:number-rows-repeated="1048000"><table:table-cell/></table:table-row>'
        '</table:table></office:spreadsheet></office:body></office:document-content>'
    )
    puffer = BytesIO()
    with zipfile.ZipFile(puffer, 'w') as archiv:
        archiv.writestr('mimetype', 'application/vnd.oasis.opendocument.spreadsheet')
        archiv.writestr('content.xml', inhalt)
    return puffer.getvalue()


def sls_datei():
    return ods_datei([
        [None, 'Klassenstufe', 1, 'LQ.Ende.1'],
        [None, 'Zeipunkt', 'Ende'],
        [],
        [None, 'Name', 'Vorname w', 'Vorname m', 'RW', 'LQ', 'Lesetand', 'Leseleistung'],
        [None, None, None, None, None, '#N/A', '#N/A', '#N/A'],
        [None, 'Abt', 'Anna', None, 17, 96, 'Ende Klasse 1', 'durchschnittlich'],
        [None, None, None, None, 9, 82, 'Vor Ende Klasse 1', 'unterdurchschnittlich'],
        [None, 'Chen', None, 'Cem', 5, 75, 'Vor Ende Klasse 1', 'schwach'],
    ])


# ----------------------------------------------------------------------
# Lesen
# ----------------------------------------------------------------------

class ImportLesenTestCase(unittest.TestCase):
    def test_hsp_workbook(self):
        datei = lese_import('Auswertung HSP - Ende 3.xlsx', hsp_mappe())
        self.assertEqual(('hsp', 3, 'ende'), (datei.format, datei.jahrgang, datei.halbjahr))
        self.assertEqual(['Wörter richtig', 'Graphemtreffer', 'Alphabetische Strategie',
                          'Orthografische Strategie', 'Morphematische Strategie'], datei.kennwerte)
        # Platzhalter und Schnittzeile fallen weg.
        self.assertEqual(['Anna Abt', 'Ben Bär', 'Zora Zett'], [z.name for z in datei.zeilen])
        anna = datei.zeilen[0]
        self.assertEqual({'rohwert': 180, 'prozentrang': 35, 't_wert': 46}, anna.werte['Graphemtreffer'])
        self.assertEqual({'rohwert': 8, 'prozentrang': 20, 't_wert': 42}, anna.werte['Morphematische Strategie'])

    def test_elfe_workbook_rounds_percentiles_and_skips_ut(self):
        datei = lese_import('Auswertungstabelle ELFE II.xlsx', elfe_mappe())
        self.assertEqual(('elfe', 3, 'ende'), (datei.format, datei.jahrgang, datei.halbjahr))
        anna = datei.zeilen[0]
        self.assertEqual({'rohwert': 40, 't_wert': 45, 'prozentrang': 31}, anna.werte['Wortverständnis'])
        self.assertEqual({'t_wert': 36, 'prozentrang': 1}, anna.werte['Gesamt'])
        self.assertTrue(any('UT' in hinweis for hinweis in datei.hinweise))

    def test_sls_ods_with_gendered_first_name_columns(self):
        datei = lese_import('Auswertungstabelle SLS.ods', sls_datei())
        self.assertEqual(('sls', 1, 'ende'), (datei.format, datei.jahrgang, datei.halbjahr))
        self.assertEqual(['Anna Abt', 'Cem Chen'], [z.name for z in datei.zeilen])
        self.assertEqual({'rohwert': 17, 'lesequotient': 96}, datei.zeilen[0].werte['Leseleistung'])
        self.assertEqual(1, datei.ohne_namen)

    def test_workbook_without_names_explains_saving(self):
        mappe = Workbook()
        blatt = mappe.active
        blatt.title = 'Gesamt'
        for spalte, wert in enumerate(['Nr.', 'Vorname', 'Name', 'W', 'PR', 'T', 'GT'], start=2):
            blatt.cell(3, spalte, wert)
        blatt['C4'] = '=S01!$G$2'  # Formel ohne gespeichertes Ergebnis
        puffer = BytesIO()
        mappe.save(puffer)
        datei = lese_import('HSP Ende 2.xlsx', puffer.getvalue())
        self.assertEqual([], datei.zeilen)
        self.assertTrue(any('öffnen und speichern' in h for h in datei.hinweise))

    def test_unknown_and_broken_files_are_rejected(self):
        leer = Workbook()
        puffer = BytesIO()
        leer.save(puffer)
        with self.assertRaisesRegex(ImportFehler, 'Format wurde nicht erkannt'):
            lese_import('irgendwas.xlsx', puffer.getvalue())
        with self.assertRaises(ImportFehler):
            lese_import('kaputt.xlsx', b'kein zip')
        with self.assertRaisesRegex(ImportFehler, 'XLSX- oder ODS'):
            lese_import('liste.csv', b'a;b')

    def test_numbers_and_names(self):
        self.assertEqual(1, zahl(0.6))
        self.assertEqual(3, zahl('2,5'))
        self.assertIsNone(zahl('#N/A'))
        self.assertIsNone(zahl(''))
        self.assertEqual(normalisiere_name('Bär'), normalisiere_name('Baer'))
        self.assertEqual(normalisiere_name('  Zoë  '), normalisiere_name('zoe'))


# ----------------------------------------------------------------------
# Oberflaeche und Speichern
# ----------------------------------------------------------------------

class ImportAblaufTestCase(unittest.TestCase):
    def setUp(self):
        self.tmpdir = tempfile.mkdtemp(prefix='km_diagnostik_import_')
        database_uri = os.environ.get('TEST_DATABASE_URL') or f"sqlite:///{self.tmpdir}/test.db"
        self.app = create_app({
            'TESTING': True, 'SECRET_KEY': 'test', 'SQLALCHEMY_DATABASE_URI': database_uri,
            'UPLOAD_FOLDER': os.path.join(self.tmpdir, 'uploads'),
            'PROTECTED_UPLOAD_FOLDER': os.path.join(self.tmpdir, 'protected_uploads'),
        })
        self.client = self.app.test_client()
        with self.app.app_context():
            db.create_all()
            klara = User(username='klara', password_hash=generate_password_hash('pass'), role='teacher')
            db.session.add_all([klara, SystemKonfiguration(schuljahr='2025/2026')])
            db.session.flush()
            db.session.add(UserKlassenzuordnung(user_id=klara.id, klasse='3a', rolle='klassenleitung'))
            ensure_klasse('3a')
            lege_vorbelegung_an()
            self.kinder = {}
            for vorname, nachname in (('Anna', 'Abt'), ('Ben', 'Baer'), ('Cem', 'Chen')):
                kind = Schueler(vorname=vorname, nachname=nachname, klasse='3a', jahrgang=3)
                db.session.add(kind)
                db.session.flush()
                self.kinder[vorname] = kind.id
            db.session.commit()
        page = self.client.get('/login')
        token = CSRF_RE.search(page.get_data(as_text=True)).group(1)
        self.client.post('/login', data={'username': 'klara', 'password': 'pass', '_csrf_token': token})

    def tearDown(self):
        with self.app.app_context():
            db.session.remove()
            db.drop_all()
        shutil.rmtree(self.tmpdir, ignore_errors=True)

    def _token(self):
        return CSRF_RE.search(self.client.get('/diagnostik/import').get_data(as_text=True)).group(1)

    def _hochladen(self, inhalt, dateiname, **extra):
        daten = {'_csrf_token': self._token(), 'aktion': 'hochladen', 'klasse': '3a',
                 'schuljahr': '2025/2026', 'datei': (BytesIO(inhalt), dateiname)}
        daten.update(extra)
        return self.client.post('/diagnostik/import', data=daten, content_type='multipart/form-data')

    @staticmethod
    def _formular(html):
        """Die Felder der Vorschau, wie der Browser sie abschicken wuerde."""
        daten = {}
        for name, wert in re.findall(r'<input type="hidden" name="([^"]+)" value="([^"]*)"', html):
            daten[name] = wert.replace('&#34;', '"').replace('&quot;', '"').replace('&amp;', '&').replace('&lt;', '<').replace('&gt;', '>').replace('&#39;', "'")
        for select in re.finditer(r'<select name="([^"]+)"[^>]*>(.*?)</select>', html, re.S):
            gewaehlt = re.search(r'<option value="([^"]*)" selected>', select.group(2))
            daten[select.group(1)] = gewaehlt.group(1) if gewaehlt else ''
        for select in re.finditer(r'<select class="form-select" id="[^"]+" name="([^"]+)">(.*?)</select>', html, re.S):
            gewaehlt = re.search(r'<option value="([^"]*)" selected>', select.group(2))
            daten[select.group(1)] = gewaehlt.group(1) if gewaehlt else ''
        daten['datum'] = re.search(r'name="datum" value="([^"]*)"', html).group(1) if 'name="datum" value=' in html else ''
        return daten

    def test_hsp_preview_matches_children_and_suggests_form(self):
        html = self._hochladen(hsp_mappe(), 'Auswertung HSP - Ende 3.xlsx').get_data(as_text=True)
        self.assertIn('HSP-Auswertungsmappe', html)
        formular = self._formular(html)
        with self.app.app_context():
            hsp3 = DiagnostikTestform.query.filter_by(name='HSP 3').one().id
        self.assertEqual(str(hsp3), formular['testform_id'])
        self.assertEqual('ende', formular['halbjahr'])
        self.assertEqual(str(self.kinder['Anna']), formular['kind_0'])
        self.assertEqual(str(self.kinder['Ben']), formular['kind_1'], 'Bär/Baer nicht zugeordnet')
        self.assertEqual('', formular['kind_2'], 'Zora gibt es in der Klasse nicht')
        self.assertIn('Morphematische Strategie ✓', html)
        self.assertIn('Ohne Ergebnis aus der Datei: Cem', html)

    def test_saving_imports_values_and_keeps_other_values(self):
        with self.app.app_context():
            testform = DiagnostikTestform.query.filter_by(name='HSP 3').one()
            wue = next(k for k in testform.kennwerte if k.name == 'Wortübergreifende Strategie')
            ergebnis = DiagnostikErgebnis(schueler_id=self.kinder['Anna'], testform_id=testform.id,
                                          schuljahr='2025/2026', halbjahr='ende', datum=date(2026, 6, 1))
            ergebnis.werte.append(DiagnostikWert(kennwert_id=wue.id, prozentrang=55))
            db.session.add(ergebnis)
            db.session.commit()

        html = self._hochladen(hsp_mappe(), 'Auswertung HSP - Ende 3.xlsx').get_data(as_text=True)
        self.assertIn('aktualisiert vorhandenes', html)
        formular = self._formular(html)
        formular.update({'_csrf_token': self._token(), 'aktion': 'speichern', 'datum': ''})
        response = self.client.post('/diagnostik/import', data=formular, follow_redirects=True)
        self.assertIn('Bitte das Datum der Durchführung angeben', response.get_data(as_text=True))
        with self.app.app_context():
            self.assertEqual(1, DiagnostikErgebnis.query.count(), 'ohne Datum gespeichert')

        formular.update({'_csrf_token': self._token(), 'datum': '2026-06-15', 'datum_1': '2026-06-18'})
        response = self.client.post('/diagnostik/import', data=formular, follow_redirects=True)
        text = response.get_data(as_text=True)
        self.assertIn('2 Ergebnis(se) übernommen, davon 1 neu und 1 aktualisiert', text)

        with self.app.app_context():
            anna = DiagnostikErgebnis.query.filter_by(schueler_id=self.kinder['Anna']).one()
            werte = {w.kennwert.name: (w.rohwert, w.prozentrang, w.t_wert) for w in anna.werte}
            self.assertEqual((180, 35, 46), werte['Graphemtreffer'])
            self.assertEqual((None, 55, None), werte['Wortübergreifende Strategie'], 'Wert ohne Spalte in der Datei geloescht')
            self.assertEqual(date(2026, 6, 15), anna.datum)
            ben = DiagnostikErgebnis.query.filter_by(schueler_id=self.kinder['Ben']).one()
            self.assertEqual(date(2026, 6, 18), ben.datum, 'Abweichendes Datum der Zeile nicht uebernommen')
            self.assertEqual(3, ben.jahrgang)
            self.assertEqual(5, len(ben.werte))
            self.assertEqual(0, DiagnostikErgebnis.query.filter_by(schueler_id=self.kinder['Cem']).count())

    def test_manual_reassignment_and_duplicates(self):
        html = self._hochladen(hsp_mappe(), 'Auswertung HSP - Ende 3.xlsx').get_data(as_text=True)
        formular = self._formular(html)
        formular.update({'_csrf_token': self._token(), 'aktion': 'speichern', 'datum': '2026-06-15',
                         'kind_2': str(self.kinder['Anna'])})
        response = self.client.post('/diagnostik/import', data=formular, follow_redirects=True)
        self.assertIn('Mehrere Zeilen sind demselben Kind zugeordnet', response.get_data(as_text=True))
        with self.app.app_context():
            self.assertEqual(0, DiagnostikErgebnis.query.count())

        formular.update({'_csrf_token': self._token(), 'kind_0': '', 'kind_2': str(self.kinder['Cem'])})
        self.client.post('/diagnostik/import', data=formular)
        with self.app.app_context():
            self.assertEqual({self.kinder['Ben'], self.kinder['Cem']}, {e.schueler_id for e in DiagnostikErgebnis.query.all()})

    def test_sls_import_derives_percentile_from_lq(self):
        html = self._hochladen(sls_datei(), 'Auswertungstabelle SLS.ods', datum='2025-07-01').get_data(as_text=True)
        self.assertIn('1 Zeile(n) haben Werte, aber keinen Namen', html)
        formular = self._formular(html)
        formular.update({'_csrf_token': self._token(), 'aktion': 'speichern', 'schuljahr': '2024/2025'})
        self.client.post('/diagnostik/import', data=formular)
        with self.app.app_context():
            ergebnisse = DiagnostikErgebnis.query.all()
            self.assertEqual({self.kinder['Anna'], self.kinder['Cem']}, {e.schueler_id for e in ergebnisse})
            anna = next(e for e in ergebnisse if e.schueler_id == self.kinder['Anna'])
            self.assertEqual(('2024/2025', 'ende', date(2025, 7, 1)), (anna.schuljahr, anna.halbjahr, anna.datum))
            self.assertEqual((17, None, 96), (anna.werte[0].rohwert, anna.werte[0].prozentrang, anna.werte[0].lesequotient))

    def test_past_school_year_stores_grade_at_that_time_and_checks_date(self):
        html = self._hochladen(hsp_mappe(), 'Auswertung HSP - Ende 1.xlsx', schuljahr='2023/2024').get_data(as_text=True)
        formular = self._formular(html)
        with self.app.app_context():
            hsp1 = DiagnostikTestform.query.filter_by(name='HSP 1+').one().id
        self.assertEqual(str(hsp1), formular['testform_id'])
        self.assertIn('2023/2024 (Nachtrag)', html)

        formular.update({'_csrf_token': self._token(), 'aktion': 'speichern', 'datum': '2026-06-15'})
        response = self.client.post('/diagnostik/import', data=formular, follow_redirects=True)
        self.assertIn('liegt nicht im Schuljahr 2023/2024', response.get_data(as_text=True))

        formular.update({'_csrf_token': self._token(), 'datum': '2024-07-02'})
        self.client.post('/diagnostik/import', data=formular)
        with self.app.app_context():
            anna = DiagnostikErgebnis.query.filter_by(schueler_id=self.kinder['Anna']).one()
            # Heute Jahrgang 3 (2025/2026), zwei Jahre vorher Jahrgang 1.
            self.assertEqual(('2023/2024', 1, date(2024, 7, 2)), (anna.schuljahr, anna.jahrgang, anna.datum))

    def test_rejects_other_class_and_unknown_files(self):
        response = self.client.post('/diagnostik/import', data={
            '_csrf_token': self._token(), 'aktion': 'hochladen', 'klasse': '4b',
            'datei': (BytesIO(hsp_mappe()), 'hsp.xlsx'),
        }, content_type='multipart/form-data')
        self.assertEqual(403, response.status_code)

        html = self._hochladen(b'nichts', 'liste.txt').get_data(as_text=True)
        self.assertIn('XLSX- oder ODS-Datei', html)

    def test_tampered_payload_is_rejected(self):
        response = self.client.post('/diagnostik/import', data={
            '_csrf_token': self._token(), 'aktion': 'speichern', 'klasse': '3a', 'payload': '{kaputt',
        }, follow_redirects=True)
        self.assertIn('Vorschau ist abgelaufen oder beschädigt', response.get_data(as_text=True))


# ----------------------------------------------------------------------
# Vorbelegung nachschaerfen
# ----------------------------------------------------------------------

class VorbelegungNachschaerfenTestCase(unittest.TestCase):
    def setUp(self):
        self.tmpdir = tempfile.mkdtemp(prefix='km_diagnostik_nachschaerfen_')
        database_uri = os.environ.get('TEST_DATABASE_URL') or f"sqlite:///{self.tmpdir}/test.db"
        self.app = create_app({'TESTING': True, 'SECRET_KEY': 'test', 'SQLALCHEMY_DATABASE_URI': database_uri,
                               'UPLOAD_FOLDER': os.path.join(self.tmpdir, 'u'),
                               'PROTECTED_UPLOAD_FOLDER': os.path.join(self.tmpdir, 'p')})
        self.ctx = self.app.app_context()
        self.ctx.push()
        db.create_all()
        lege_vorbelegung_an()
        db.session.commit()

    def tearDown(self):
        db.session.remove()
        db.drop_all()
        self.ctx.pop()
        shutil.rmtree(self.tmpdir, ignore_errors=True)

    def _alte_vorbelegung(self, name):
        """Stand vor dem Abgleich: Strategien nur mit PR, keine Mitte-Testung."""
        testform = DiagnostikTestform.query.filter_by(name=name).one()
        for kennwert in testform.kennwerte:
            if 'Strategie' in kennwert.name:
                kennwert.rohwert = kennwert.t_wert = False
        for zeitpunkt in list(testform.zeitpunkte):
            if zeitpunkt.halbjahr == 'mitte':
                testform.zeitpunkte.remove(zeitpunkt)
        db.session.commit()
        return testform

    def test_fresh_preset_needs_nothing(self):
        self.assertEqual(0, schaerfe_vorbelegung_nach())
        self.assertEqual(['Mitte Klasse 2', 'Ende Klasse 2'],
                         [z.label for z in DiagnostikTestform.query.filter_by(name='HSP 2').one().zeitpunkte_sortiert])

    def test_old_preset_is_updated_unless_results_exist(self):
        hsp2 = self._alte_vorbelegung('HSP 2')
        hsp3 = self._alte_vorbelegung('HSP 3')
        kind = Schueler(vorname='A', nachname='B', klasse='3a')
        db.session.add(kind)
        db.session.flush()
        db.session.add(DiagnostikErgebnis(schueler_id=kind.id, testform_id=hsp3.id, schuljahr='2025/2026', halbjahr='ende'))
        db.session.commit()

        self.assertEqual(1, schaerfe_vorbelegung_nach())
        db.session.commit()
        self.assertEqual(0, schaerfe_vorbelegung_nach(), 'nicht idempotent')

        strategie = DiagnostikKennwert.query.filter_by(testform_id=hsp2.id, name='Alphabetische Strategie').one()
        self.assertEqual(['rohwert', 'prozentrang', 't_wert'], strategie.wertarten)
        self.assertEqual(2, DiagnostikZeitpunkt.query.filter_by(testform_id=hsp2.id).count())
        # HSP 3 hat ein Ergebnis und bleibt unberuehrt.
        self.assertEqual(1, DiagnostikZeitpunkt.query.filter_by(testform_id=hsp3.id).count())


class ZuordnungTestCase(unittest.TestCase):
    class Kind:
        def __init__(self, id, vorname, nachname):
            self.id, self.vorname, self.nachname = id, vorname, nachname

    def test_first_name_only_when_unique(self):
        from diagnostik_import import ImportZeile
        kinder = [self.Kind(1, 'Anna', 'Abt'), self.Kind(2, 'Ben', 'Bär'), self.Kind(3, 'Ben', 'Berg')]
        zeilen = [ImportZeile('', 'Anna'), ImportZeile('', 'Ben'), ImportZeile('Berg', 'Ben'), ImportZeile('Abt', 'Anna')]
        zuordnung = ordne_kinder_zu(zeilen, kinder)
        self.assertEqual(1, zuordnung[0].id)
        self.assertIsNone(zuordnung[1], 'Vorname Ben ist nicht eindeutig')
        self.assertEqual(3, zuordnung[2].id)
        self.assertIsNone(zuordnung[3], 'Anna doppelt vergeben')


if __name__ == '__main__':
    unittest.main()
